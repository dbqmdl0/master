"""산출물 정합성 검사기 — 코드·설정·워크플로 설계·데이터 구조가 서로 맞는지 기계적으로 확인한다.

실행: python3 tools/check_consistency.py   (표준 라이브러리 + 선택적으로 openpyxl)
실패 항목이 있으면 종료코드 1.
"""
from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
P1 = ROOT / "pilots/pilot_01/automation_assets"
P2 = ROOT / "pilots/pilot_02/automation_assets"
P3 = ROOT / "pilots/pilot_03/automation_assets"
SYSTEM_COLUMNS = {"ID", "Created", "Modified", "Author", "Editor"}

results: list[tuple[bool, str]] = []


def check(ok: bool, msg: str):
    results.append((bool(ok), msg))


def load_json(p: Path):
    return json.loads(p.read_text(encoding="utf-8"))


def lists_index(*schema_paths: Path) -> dict[str, dict]:
    idx = {}
    for p in schema_paths:
        for lst in load_json(p)["lists"]:
            idx[lst["name"]] = {c["internalName"]: c for c in lst["columns"]}
    return idx


def check_all_json():
    files = [p for p in ROOT.rglob("*.json") if "build" not in p.parts and "node_modules" not in p.parts]
    bad = []
    for p in files:
        try:
            load_json(p)
        except Exception as e:  # noqa: BLE001
            bad.append(f"{p.relative_to(ROOT)}: {e}")
    check(not bad, f"JSON 파일 {len(files)}개 구문 유효" + (f" — 오류: {bad}" if bad else ""))


def check_flow_refs(spec_path: Path, lists: dict[str, dict]):
    spec = load_json(spec_path)
    missing = []
    for f in spec["flows"]:
        for kind in ("reads", "writes"):
            for lst, cols in f.get(kind, {}).items():
                if lst not in lists:
                    missing.append(f"{f['id']}: 목록 {lst} 없음")
                    continue
                for c in cols:
                    if c not in lists[lst] and c not in SYSTEM_COLUMNS:
                        missing.append(f"{f['id']}.{kind}: {lst}.{c}")
        for cond in f.get("trigger", {}).get("triggerConditions", []):
            for col in re.findall(r"body/([A-Za-z]+)", cond):
                tl = f["trigger"].get("list")
                if tl and col not in lists.get(tl, {}) and col not in SYSTEM_COLUMNS:
                    missing.append(f"{f['id']} 트리거 조건: {tl}.{col}")
        if "inputCard" in f.get("trigger", {}):
            card = spec_path.parent / f["trigger"]["inputCard"]
            if not card.exists():
                missing.append(f"{f['id']}: 카드 파일 없음 {card.name}")
    check(not missing, f"{spec_path.relative_to(ROOT)} 열·목록·카드 참조 일치" + (f" — 불일치: {missing}" if missing else ""))


def choices(lists, lst, col):
    return lists[lst][col]["choices"]


def check_pilot01(lists):
    sys.path.insert(0, str(P1))
    import request_hub_rules as R  # noqa: E402
    check(choices(lists, "RequestHub", "Status") == R.STATUSES, "P1 상태 선택지 == request_hub_rules.STATUSES")
    md = (P1 / "expressions.md").read_text(encoding="utf-8")
    m = re.search(r"```json\n(\{.*?\})\n```", md, re.S)
    tr = json.loads(m.group(1)) if m else {}
    same = tr and all(set(tr[k]) == R.TRANSITIONS[k] for k in R.TRANSITIONS) and set(tr) == set(R.TRANSITIONS)
    check(same, "P1 expressions.md 전이표 == request_hub_rules.TRANSITIONS")
    with (P1 / "templates/SLAPolicy_seed.csv").open(encoding="utf-8") as f:
        policy_types = [r["Title"] for r in csv.DictReader(f)]
    check(policy_types == choices(lists, "RequestHub", "RequestType"), "P1 SLAPolicy 시드 == 요청유형 선택지")
    card = load_json(P1 / "cards/capture_request_from_message.json")
    by_id = {b.get("id"): b for b in card["body"] if "id" in b}
    check([c["value"] for c in by_id["requestType"]["choices"]] == choices(lists, "RequestHub", "RequestType"),
          "P1 Teams 등록 카드 요청유형 == List 선택지")
    check([c["value"] for c in by_id["requestingDept"]["choices"]] == choices(lists, "RequestHub", "RequestingDept"),
          "P1 Teams 등록 카드 요청부서 == List 선택지")
    with (P1 / "templates/RequestHub_import_sample.csv").open(encoding="utf-8") as f:
        hdr = next(csv.reader(f))
    check(all(h in lists["RequestHub"] for h in hdr), "P1 가져오기 샘플 CSV 열 ⊂ RequestHub 스키마")
    rows = list(csv.DictReader((P1 / "templates/RequestHub_import_sample.csv").open(encoding="utf-8")))
    check(all(r["Status"] in R.STATUSES and r["RequestType"] in policy_types for r in rows), "P1 샘플 행 값이 선택지 안에 있음")


def ts_const(src: str, name: str) -> list[str]:
    m = re.search(rf"const {name}: string\[\] = \[(.*?)\];", src, re.S)
    return re.findall(r'"([^"]+)"', m.group(1)) if m else []


def ts_main_params(src: str) -> list[str]:
    m = re.search(r"function main\(workbook: ExcelScript\.Workbook,?(.*?)\)\s*:", src, re.S)
    return re.findall(r"(\w+)\s*:", m.group(1)) if m else []


def check_pilot02():
    contract = load_json(P2 / "data_contract.json")
    cols = [c["name"] for c in contract["columns"]]
    req = [c["name"] for c in contract["columns"] if c["required"]]
    v = (P2 / "office_scripts/ValidateSubmission.ts").read_text(encoding="utf-8")
    a = (P2 / "office_scripts/AppendToMaster.ts").read_text(encoding="utf-8")
    x = (P2 / "office_scripts/BuildExceptionReport.ts").read_text(encoding="utf-8")
    check(ts_const(v, "COLUMNS") == cols, "P2 ValidateSubmission.COLUMNS == 데이터 계약")
    check(ts_const(v, "REQUIRED") == req, "P2 ValidateSubmission.REQUIRED == 계약 필수열")
    check(ts_const(a, "CONTRACT_COLUMNS") == cols and ts_const(a, "SYSTEM_COLUMNS") == contract["systemColumns"],
          "P2 AppendToMaster 열 정의 == 데이터 계약")
    codes = {r["code"] for r in contract["rules"]}
    check(codes == set(re.findall(r'"([EW]\d\d)"', v)), "P2 계약 규칙 코드 == ValidateSubmission 구현 코드")
    xcodes = {r["code"] for r in contract["exceptionRules"]}
    check(xcodes == set(re.findall(r'"(X\d\d)"', x)), "P2 예외 규칙 코드 == BuildExceptionReport 구현 코드")
    spec = load_json(P2 / "flows_spec.json")
    srcs = {"ValidateSubmission": v, "AppendToMaster": a, "BuildExceptionReport": x}
    bad = [s["script"] for f in spec["flows"] for s in f.get("scripts", [])
           if ts_main_params(srcs[s["script"]]) != s["params"]]
    check(not bad, "P2 Flow 명세 스크립트 매개변수 == main() 시그니처" + (f" — 불일치: {bad}" if bad else ""))
    try:
        from openpyxl import load_workbook
        wb = load_workbook(P2 / "templates/PlantSubmission_Template.xlsx")
        ws = wb["Submit"]
        hdr = [c.value for c in ws[1]][: len(cols)]
        check(hdr == cols and "tblSubmit" in ws.tables, "P2 제출 템플릿 머리글·표 이름 == 계약")
        mb = load_workbook(P2 / "templates/SCM_Master_Fact.xlsx")
        fh = [c.value for c in mb["Fact"][1]]
        check(fh == cols + contract["systemColumns"] and "tblFact" in mb["Fact"].tables, "P2 마스터 tblFact 머리글 == 계약+시스템열")
    except ImportError:
        check(True, "P2 템플릿 검사 생략(openpyxl 없음)")


def check_pilot03(lists):
    sys.path.insert(0, str(P3 / "python"))
    import daily_brief as B  # noqa: E402
    check(B.OPEN_ACTION <= set(choices(lists, "ActionItems", "Status")), "P3 브리핑 미결 상태 ⊂ ActionItems 상태 선택지")
    check("결정요청" in choices(lists, "DecisionLog", "Status"), "P3 DecisionLog에 '결정요청' 상태 존재")
    check("회신대기" in choices(lists, "ActionItems", "ActionType"), "P3 ActionItems에 '회신대기' 유형 존재")
    md = (P3 / "agent/decision_assistant_instructions.md").read_text(encoding="utf-8")
    block = re.search(r"```text\n(.*?)```", md, re.S).group(1)
    instr = (P3 / "agent/instructions.txt").read_text(encoding="utf-8")
    check(block == instr, "P3 에이전트 instructions.txt == 설정서 본문")
    check(len(instr) <= 8000, f"P3 에이전트 지침 8,000자 이내({len(instr)}자)")
    manifest = load_json(P3 / "agent/declarativeAgent.json")
    ref = re.search(r"file\('([^']+)'\)", manifest["instructions"]).group(1)
    check((P3 / "agent" / ref).exists(), "P3 매니페스트가 참조하는 지침 파일 존재")
    ids = {r["DecisionID"] for r in csv.DictReader((P3 / "sample_data/DecisionLog_mock.csv").open(encoding="utf-8"))}
    ev = load_json(P3 / "agent/eval_cases.json")
    refs = {t for c in ev["cases"] for t in c["expect"].get("mustInclude", []) + c["expect"].get("mustNotInclude", [])
            if t.startswith("DEC-")}
    check(refs <= ids, "P3 평가셋 결정번호가 모의 DecisionLog에 존재")
    check(sum(1 for c in ev["cases"] if c["type"] == "unknown") == ev["passCriteria"]["unknownCorrect"],
          "P3 평가셋 '모름' 문항 수 == 합격 기준")


def check_matrix():
    rows = list(csv.reader((ROOT / "04_Prioritization_Matrix.csv").open(encoding="utf-8-sig")))
    data = [r for r in rows[1:] if r and r[0] and not r[0].startswith("#")]
    check(len(data) >= 30, f"우선순위 매트릭스 과제 {len(data)}개(≥30)")
    top = [r for r in data if r[23].startswith("TOP10-")]
    check(len(top) == 10, "TOP10 정확히 10개")
    approval = [r for r in data if r[22] == "Y"]
    check(all(r[0] == "별도승인" for r in approval), "고위험 과제는 순위 대신 '별도승인'으로 분리")


def main() -> int:
    lists = lists_index(P1 / "lists_schema.json", P3 / "lists_schema.json")
    check_all_json()
    check_flow_refs(P1 / "flows_spec.json", lists)
    check_flow_refs(P3 / "flows_spec.json", lists)
    check_pilot01(lists)
    check_pilot02()
    check_pilot03(lists)
    check_matrix()
    for ok, msg in results:
        print(("PASS " if ok else "FAIL ") + msg)
    failed = sum(1 for ok, _ in results if not ok)
    print(f"\n{len(results) - failed}/{len(results)} 통과")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())

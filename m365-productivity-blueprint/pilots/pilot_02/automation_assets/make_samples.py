"""Pilot 02 모의 데이터·Excel 템플릿 생성기 (가상 데이터만 생성).

실행: python3 make_samples.py
생성물
  sample_data/  : 공장 제출 CSV(정상/오류 주입), SKU 마스터, Python 검증 기대결과(교차검증용)
  templates/    : PlantSubmission_Template.xlsx(tblSubmit), SCM_Master_Fact.xlsx(tblFact·tblSKU, 모의 이력)
openpyxl 이 없으면 CSV만 만들고 템플릿은 건너뛴다.
"""
from __future__ import annotations

import csv
import json
import random
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE / "python"))
import psi_validator as pv  # noqa: E402

CONTRACT = pv.load_contract()
COLS = [c["name"] for c in CONTRACT["columns"]]
SYS = CONTRACT["systemColumns"]
SKUS = [f"SKU-MOCK-{i:03d}" for i in range(101, 113)]
LINES = ["L1", "L2"]
random.seed(42)


def make_rows(week: str, plant: str, begin: dict | None = None) -> list[list]:
    rows = []
    for i, sku in enumerate(SKUS[:8]):
        line = LINES[i % 2]
        b = (begin or {}).get((sku, line), random.randint(500, 3000))
        prod = random.randint(300, 1200)
        rec = random.choice([0, 0, 100])
        ship = random.randint(300, 1200)
        end = max(0, b + prod + rec - ship)
        adj = end - (b + prod + rec - ship)  # 음수 방지분은 조정으로 기록
        rows.append([week, plant, sku, line, b, prod, rec, ship, adj, end,
                     random.randint(300, 1200), random.randint(30, 40), 40,
                     "재고 실사 차이" if adj else ""])
    return rows


def write_csv(path: Path, rows: list[list]):
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(COLS)
        w.writerows(rows)


def main() -> int:
    sd = HERE / "sample_data"
    sd.mkdir(exist_ok=True)
    (sd / "sku_master.txt").write_text("\n".join(SKUS) + "\n", encoding="utf-8")

    ok_rows = make_rows("2026-W44", "PLT-A")
    write_csv(sd / "PLT-A_2026-W44_ok.csv", ok_rows)

    bad = make_rows("2026-W44", "PLT-B")
    bad[0][0] = "2026-W43"                      # E03
    bad[1][2] = "SKU-UNKNOWN"                   # E05
    bad[2][7] = "1,200"                         # E06 (천단위 쉼표 문자열)
    bad[3][9] = int(bad[3][9]) + 50             # E08
    bad[4][4] = ""                              # E02 (+E08)
    bad.append(list(bad[5]))                    # E07 (6행과 같은 키)
    r = bad[6]                                  # W01: 조정 −10, 비고 없음, 수지는 맞춤
    r[4], r[5], r[6], r[7], r[8], r[13] = 2000, 500, 0, 400, -10, ""
    r[9] = r[4] + r[5] + r[6] - r[7] + r[8]
    bad[7][11] = 45                             # W02: 가동 45h > 가용 40h
    write_csv(sd / "PLT-B_2026-W44_errors.csv", bad)

    expected = {}
    for name, plant in (("PLT-A_2026-W44_ok.csv", "PLT-A"), ("PLT-B_2026-W44_errors.csv", "PLT-B")):
        header, body = pv.read_csv(sd / name)
        r = pv.validate(header, body, "2026-W44", plant, SKUS, CONTRACT)
        expected[name] = sorted(f"{i['code']}@{i['row']}" for i in r["issues"])
    (sd / "expected_issues.json").write_text(json.dumps(expected, ensure_ascii=False, indent=2), encoding="utf-8")

    try:
        from openpyxl import Workbook
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        print("openpyxl 없음: 템플릿 생략")
        return 0

    tp = HERE / "templates"
    tp.mkdir(exist_ok=True)
    labels = {c["name"]: c["label"] for c in CONTRACT["columns"]}

    wb = Workbook()
    guide = wb.active
    guide.title = "작성안내"
    guide.append(["공장 주간 제출 템플릿 (데이터 계약 v" + CONTRACT["contractVersion"] + ") — 가상 예시"])
    guide.append(["1. 'Submit' 시트의 표(tblSubmit)에만 입력하고, 열 이름·순서를 바꾸지 마세요."])
    guide.append(["2. 한 행 = 주차 × 공장 × SKU × 라인. 숫자 칸에는 숫자만(쉼표·단위 금지)."])
    guide.append(["3. 기말재고 = 기초 + 생산 + 이동입고 − 출하 + 조정. 조정이 있으면 비고에 사유."])
    guide.append(["4. 저장 위치: SCM 제출 라이브러리 / <주차> / <공장코드> 폴더. 저장 즉시 자동 검증 결과가 Teams로 옵니다."])
    guide.append([])
    guide.append(["열", "의미", "필수", "규칙"])
    for c in CONTRACT["columns"]:
        rule = c.get("pattern") or ("0 이상" if c.get("min") == 0 else "")
        guide.append([c["name"], c["label"], "Y" if c["required"] else "", rule])
    ws = wb.create_sheet("Submit")
    ws.append(COLS)
    ws.append([""] * len(COLS))
    t = Table(displayName="tblSubmit", ref=f"A1:{chr(64 + len(COLS))}2")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        error="0 이상의 숫자만 입력하세요", showErrorMessage=True)
    ws.add_data_validation(dv)
    for i, c in enumerate(CONTRACT["columns"]):
        col = chr(65 + i)
        ws.column_dimensions[col].width = 13
        ws.cell(row=1, column=i + 1).comment = None
        if c["type"] == "number" and c.get("min") == 0:
            dv.add(f"{col}2:{col}5000")
    wb.save(tp / "PlantSubmission_Template.xlsx")

    mb = Workbook()
    fs = mb.active
    fs.title = "Fact"
    fs.append(COLS + SYS)
    hist = []
    for plant in ("PLT-A", "PLT-B"):
        begin = None
        for w in ("2026-W41", "2026-W42", "2026-W43"):
            rows = make_rows(w, plant, begin)
            begin = {(r[2], r[3]): r[9] for r in rows}
            hist += [r + [f"{w}|{plant}", "2026-10-01T00:00:00Z", "mock"] for r in rows]
    for r in hist:
        fs.append(r)
    ft = Table(displayName="tblFact", ref=f"A1:{chr(64 + len(COLS) + len(SYS))}{len(hist) + 1}")
    ft.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    fs.add_table(ft)
    ss = mb.create_sheet("SKU")
    ss.append(["SKU", "Active"])
    for s in SKUS:
        ss.append([s, "Y"])
    st = Table(displayName="tblSKU", ref=f"A1:B{len(SKUS) + 1}")
    ss.add_table(st)
    mb.create_sheet("Exceptions")
    mb.save(tp / "SCM_Master_Fact.xlsx")
    print("samples + templates written")
    return 0


if __name__ == "__main__":
    sys.exit(main())
